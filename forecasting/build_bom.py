"""Build data/outputs/forecast/ui/bom_AIL.json from the AIL BOM workbook.

Usage: python forecasting/build_bom.py
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw/AIL_Steel_Raw_Material_BOM_Oct2026.xlsx"
OUT = ROOT / "data/outputs/forecast/ui/bom_AIL.json"
AS_OF = "2026-10-12"


def num(v):
    return float(v) if isinstance(v, (int, float)) else None


def slug(name):
    return re.sub(r"[^A-Z0-9]", "", str(name).upper())


def main():
    wb = openpyxl.load_workbook(RAW, data_only=True)
    mix, dem = wb["Steel Mix BOM"], wb["Material Demand Forecast"]
    rows = list(mix.iter_rows(values_only=True))

    hi = next(i for i, r in enumerate(rows) if r[0] == "Material" and str(r[1] or "").startswith("Consumed"))
    materials, charge, cost = [], None, None
    for r in rows[hi + 1:]:
        if r[0] is None or not str(r[0]).strip():
            continue
        lab = str(r[0]).strip()
        if lab == "Total charge":
            charge, cost = num(r[6]), num(r[9])
            break
        if lab.startswith(("Metallic yield", "Steel needed per MT")):
            continue
        materials.append({
            "id": slug(lab),
            "name": lab,
            "mix_pct": num(r[2]) or 0.0,
            "kg_per_mt_steel": num(r[6]) or 0.0,
            "rate": num(r[8]) or 0.0,
            "stock_mt": 0.0,
        })

    yld = next((num(r[1]) for r in rows if str(r[0] or "").startswith("Metallic yield")), None)
    spb = next((num(r[1]) for r in rows if str(r[0] or "").startswith("Steel needed per MT")), None)

    energy = []
    ei = next((i for i, r in enumerate(rows) if str(r[0] or "").strip() == "Item" and str(r[1] or "").strip() == "Norm"), None)
    if ei is not None:
        for r in rows[ei + 1:]:
            if r[0] is None or not str(r[0]).strip():
                break
            energy.append({"name": str(r[0]).strip(), "value": num(r[1]), "unit": str(r[2] or "").strip()})

    drows = list(dem.iter_rows(values_only=True))
    si = next(i for i, r in enumerate(drows) if r[0] == "Material" and str(r[6] or "").startswith("Stock"))
    stock = {}
    for r in drows[si + 1:]:
        if r[0] is None or not str(r[0]).strip() or str(r[0]).strip() == "Total":
            break
        stock[str(r[0]).strip()] = num(r[6]) or 0.0
    for m in materials:
        m["stock_mt"] = round(stock.get(m["name"], 0.0), 3)

    payload = {
        "sbu": "AIL",
        "as_of": AS_OF,
        "currency": "BDT",
        "steel_per_bar": round(spb, 6),
        "metallic_yield": round(yld, 6),
        "charge_kg_per_mt_steel": round(charge, 3),
        "material_cost_per_mt_steel": round(cost, 2),
        "materials": materials,
        "energy": energy,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(f"wrote {OUT} ({len(materials)} materials, {len(energy)} energy norms)")


if __name__ == "__main__":
    main()
